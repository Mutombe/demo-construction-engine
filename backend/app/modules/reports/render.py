import csv
import io
from datetime import date
from dataclasses import dataclass, field
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ColumnKind = Literal["text", "money", "qty", "date", "int", "pct"]


@dataclass
class ReportColumn:
    key: str
    label: str
    kind: ColumnKind = "text"
    width: int = 14


@dataclass
class ReportTable:
    title: str
    columns: list[ReportColumn]
    rows: list[dict] = field(default_factory=list)
    totals: dict | None = None  # keyed subset of column keys


_NUMBER_FORMATS: dict[ColumnKind, str] = {
    "money": "#,##0.00",
    "qty": "#,##0.000",
    "int": "#,##0",
    "pct": "0.0%",
    "date": "yyyy-mm-dd",
    "text": "@",
}


def to_csv(table: ReportTable) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([c.label for c in table.columns])
    for row in table.rows:
        writer.writerow([row.get(c.key, "") for c in table.columns])
    if table.totals:
        writer.writerow(
            ["TOTAL" if i == 0 else table.totals.get(c.key, "")
             for i, c in enumerate(table.columns)]
        )
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(table: ReportTable) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:31]  # Excel sheet name limit

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2A4B8D", end_color="2A4B8D", fill_type="solid")
    for col_index, column in enumerate(table.columns, start=1):
        cell = ws.cell(row=1, column=col_index, value=column.label)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_index)].width = column.width

    for row_index, row in enumerate(table.rows, start=2):
        for col_index, column in enumerate(table.columns, start=1):
            value = row.get(column.key)
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.number_format = _NUMBER_FORMATS[column.kind]

    if table.totals:
        total_row = len(table.rows) + 2
        bold = Font(bold=True)
        first = ws.cell(row=total_row, column=1, value="TOTAL")
        first.font = bold
        for col_index, column in enumerate(table.columns, start=1):
            if column.key in table.totals:
                cell = ws.cell(row=total_row, column=col_index, value=table.totals[column.key])
                cell.font = bold
                cell.number_format = _NUMBER_FORMATS[column.kind]

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(table: ReportTable, company_name: str = "Company Name Not Set") -> bytes:
    """The same table, for someone who is going to print it or send it on.

    Landscape, because a report built for a spreadsheet has more columns than
    a portrait page can carry without shrinking the type past reading size.
    """
    from app.common.pdf import Col, DocumentPdf

    doc = DocumentPdf(
        company_name=company_name,
        doc_title=table.title,
        doc_number=date.today().isoformat(),
        orientation="L",
    )

    # Widths come from the spreadsheet hints, normalised to the page: a column
    # someone sized for Excel is the best available signal of what matters.
    total_width = sum(column.width for column in table.columns) or 1
    cols = [
        Col(
            column.label,
            column.width / total_width,
            "R" if column.kind in ("money", "qty", "int", "pct") else "L",
        )
        for column in table.columns
    ]

    rows = [
        [_pdf_value(row.get(column.key), column.kind) for column in table.columns]
        for row in table.rows
    ]
    totals = None
    if table.totals:
        totals = [
            _pdf_value(table.totals.get(column.key), column.kind)
            if column.key in table.totals
            else ("TOTAL" if column is table.columns[0] else "")
            for column in table.columns
        ]

    if rows:
        doc.table(cols=cols, rows=rows, totals=totals)
    else:
        doc.note("Nothing to report for this selection.")
    return doc.render()


def _pdf_value(value, kind: ColumnKind) -> str:
    from app.common.pdf import money

    if value is None or value == "":
        return "-"
    if kind == "money":
        return money(value)
    if kind == "pct":
        return f"{float(value) * 100:,.1f}%"
    if kind == "qty":
        return f"{float(value):,.3f}"
    if kind == "int":
        return f"{int(value):,}"
    return str(value)
