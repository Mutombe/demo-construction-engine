import io
from decimal import Decimal

from openpyxl import load_workbook

from app.common.enums import UserRole
from tests.factories import (
    auth_headers,
    make_boq_item,
    make_boq_section,
    make_cost_entry,
    make_project,
    make_user,
)


def _viewer(db):
    return auth_headers(make_user(db, role=UserRole.viewer))


def test_project_cost_csv(client, db):
    headers = _viewer(db)
    project = make_project(db)
    section = make_boq_section(db, project, code="B", title="Substructure")
    item = make_boq_item(
        db, section, description="Concrete G25", quantity=Decimal("10"), rate=Decimal("100")
    )
    make_cost_entry(db, project, boq_item_id=item.id, amount=Decimal("650"))
    make_cost_entry(db, project, amount=Decimal("120"))  # unallocated

    res = client.get(
        f"/api/v1/reports/project-cost?project_id={project.id}&format=csv", headers=headers
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment" in res.headers["content-disposition"]
    text = res.content.decode("utf-8-sig")
    assert "Concrete G25" in text
    assert "Unallocated costs" in text
    assert text.strip().splitlines()[-1].startswith("TOTAL")
    assert "650" in text and "120" in text


def test_project_cost_xlsx_opens(client, db):
    headers = _viewer(db)
    project = make_project(db)
    section = make_boq_section(db, project)
    make_boq_item(db, section)

    res = client.get(
        f"/api/v1/reports/project-cost?project_id={project.id}&format=xlsx", headers=headers
    )
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header[:3] == ["Section", "Item", "Description"]


def test_cost_ledger_filters(client, db):
    headers = _viewer(db)
    project = make_project(db)
    make_cost_entry(db, project, amount=Decimal("100"), source="manual")
    make_cost_entry(db, project, amount=Decimal("200"), source="expense")

    res = client.get(
        f"/api/v1/reports/cost-ledger?project_id={project.id}&source=expense", headers=headers
    )
    text = res.content.decode("utf-8-sig")
    assert "expense" in text
    assert "manual" not in text
    assert "200" in text


def test_cost_ledger_requires_project(client, db):
    headers = _viewer(db)
    res = client.get("/api/v1/reports/cost-ledger", headers=headers)
    assert res.status_code == 422


def test_stock_valuation_math(client, db):
    proc = auth_headers(make_user(db, role=UserRole.procurement_officer))
    item = client.post(
        "/api/v1/stock-items",
        json={"code": "RPT-1", "name": "Report item", "unit": "ea"},
        headers=proc,
    ).json()
    client.post(
        f"/api/v1/stock-items/{item['id']}/goods-in",
        json={"quantity": "10", "unit_cost": "25.00"},
        headers=proc,
    )
    res = client.get("/api/v1/reports/stock-valuation?format=csv", headers=proc)
    text = res.content.decode("utf-8-sig")
    assert "RPT-1" in text
    assert "250.00" in text  # value = 10 × 25


def test_procurement_register_includes_all_types(client, db):
    from tests.factories import make_quote, make_rfq, make_supplier

    headers = _viewer(db)
    project = make_project(db)
    section = make_boq_section(db, project)
    boq_item = make_boq_item(db, section)
    rfq = make_rfq(db, project, [boq_item], doc_number="RFQ-RPT-001")
    supplier = make_supplier(db, name="Register Supplies")
    make_quote(db, rfq, supplier)

    res = client.get("/api/v1/reports/procurement-register", headers=headers)
    text = res.content.decode("utf-8-sig")
    assert "RFQ-RPT-001" in text
    assert "Register Supplies" in text
    assert "RFQ" in text and "Quote" in text


def test_reports_require_auth(client, db):
    res = client.get("/api/v1/reports/stock-valuation")
    assert res.status_code == 401
